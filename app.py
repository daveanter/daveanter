"""
Employee Shift Schedule System
Flask web app — employee form, manager approval, Excel export with visual timeline.

Usage:
  python app.py serve                          Start web server
  python app.py send-invites --week YYYY-MM-DD Email schedule requests to all employees
"""

from flask import (Flask, render_template, request, redirect,
                   url_for, session, abort, send_file, jsonify)
import sqlite3, secrets, smtplib, socket, argparse, sys, os
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from datetime import date, timedelta
from functools import wraps
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', secrets.token_hex(32))

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
DB_PATH    = os.path.join(BASE_DIR, 'schedule.db')
CONFIG_PATH = os.path.join(BASE_DIR, 'config.json')

DAYS       = ['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday','Sunday']
DAYS_SHORT = ['Mon','Tue','Wed','Thu','Fri','Sat','Sun']
WEEKDAYS   = DAYS[:5]
WEEKDAYS_SHORT = DAYS_SHORT[:5]

LUNCH_OPTIONS = [(0,'No lunch'),(15,'15 min'),(30,'30 min'),
                 (45,'45 min'),(60,'1 hour'),(90,'1.5 hrs'),(120,'2 hrs')]

# 8:00 AM → 9:00 PM in 30-min increments (26 slots)
TIMELINE_SLOTS = list(range(8*60, 21*60, 30))  # minutes since midnight

STATUS_META = {
    'pending':           ('Pending',           'secondary'),
    'submitted':         ('Submitted',          'primary'),
    'changes_requested': ('Changes Sent',       'warning'),
    'employee_approved': ('Employee Approved',  'info'),
    'approved':          ('Approved ✓',         'success'),
}

# ─── Database ────────────────────────────────────────────────────────────────

def get_config():
    with open(CONFIG_PATH) as f:
        return json.load(f)

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS employees (
            id    INTEGER PRIMARY KEY AUTOINCREMENT,
            name  TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE
        );
        CREATE TABLE IF NOT EXISTS schedule_requests (
            id             INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id    INTEGER NOT NULL REFERENCES employees(id),
            week_start     TEXT NOT NULL,
            employee_token TEXT UNIQUE NOT NULL,
            accept_token   TEXT UNIQUE NOT NULL,
            manager_token  TEXT UNIQUE NOT NULL,
            status         TEXT NOT NULL DEFAULT 'pending',
            manager_notes  TEXT NOT NULL DEFAULT '',
            created_at     TEXT NOT NULL DEFAULT (datetime('now')),
            submitted_at   TEXT,
            approved_at    TEXT
        );
        CREATE TABLE IF NOT EXISTS shifts (
            id               INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id       INTEGER NOT NULL REFERENCES schedule_requests(id),
            day_index        INTEGER NOT NULL,
            start_time       TEXT,
            end_time         TEXT,
            lunch_start_time TEXT,
            lunch_minutes    INTEGER NOT NULL DEFAULT 30,
            is_off           INTEGER NOT NULL DEFAULT 0,
            proposed_by      TEXT NOT NULL DEFAULT 'employee'
        );
    """)
    conn.commit()
    conn.close()

# ─── Helpers ─────────────────────────────────────────────────────────────────

def to_min(t):
    """'HH:MM' → minutes since midnight, or None."""
    if not t:
        return None
    try:
        h, m = map(int, t.split(':'))
        return h * 60 + m
    except Exception:
        return None

def worked_hours(start, end, lunch_mins, is_off):
    if is_off:
        return 0.0
    s, e = to_min(start), to_min(end)
    if s is None or e is None:
        return 0.0
    return max(0.0, (e - s - (lunch_mins or 0)) / 60.0)

def lunch_start_minutes(start, end, lunch_start, lunch_mins):
    """Return lunch start in minutes since midnight."""
    if not lunch_mins:
        return None
    if lunch_start:
        return to_min(lunch_start)
    # Estimate: midpoint of working period
    s, e = to_min(start), to_min(end)
    if s is None or e is None:
        return None
    mid = (s + e) // 2
    return mid - (lunch_mins // 2)

def slot_fill(slot, s_min, e_min, ls_min, le_min):
    """Return hex color for a 30-min timeline slot, or None if unworked."""
    if s_min is None or e_min is None:
        return None
    if slot < s_min or slot >= e_min:
        return None
    if ls_min is not None and ls_min <= slot < le_min:
        return "FFEB9C"    # lunch — yellow
    return "4472C4"        # on shift — blue

def build_day_data(week_start, existing):
    """Build per-day dicts for form rendering."""
    days = []
    for i in range(7):
        s = existing.get(i, {})
        days.append({
            'index':           i,
            'name':            DAYS[i],
            'date':            (week_start + timedelta(days=i)).strftime('%b %d'),
            'start_time':      s.get('start_time')      or '09:00',
            'end_time':        s.get('end_time')        or '17:00',
            'lunch_start_time':s.get('lunch_start_time')or '',
            'lunch_minutes':   s.get('lunch_minutes', 30),
            'is_off':          bool(s.get('is_off', 0)),
        })
    return days

# ─── Auth ─────────────────────────────────────────────────────────────────────

def mgr_only(f):
    @wraps(f)
    def wrap(*a, **kw):
        if not session.get('is_manager'):
            return redirect(url_for('manager_login'))
        return f(*a, **kw)
    return wrap

@app.route('/admin/login', methods=['GET','POST'])
def manager_login():
    if request.method == 'POST':
        cfg = get_config()
        if request.form.get('password') == cfg.get('manager_password','changeme'):
            session['is_manager'] = True
            return redirect(url_for('manager_dashboard'))
        return render_template('manager_login.html', error='Incorrect password.')
    return render_template('manager_login.html', error=None)

@app.route('/admin/logout')
def manager_logout():
    session.clear()
    return redirect(url_for('manager_login'))

# ─── Employee routes ──────────────────────────────────────────────────────────

def _parse_shifts_from_form():
    shifts = []
    for i in range(7):
        is_off = 1 if request.form.get(f'd{i}_off') else 0
        if is_off:
            shifts.append({'day_index':i,'start_time':None,'end_time':None,
                           'lunch_start_time':None,'lunch_minutes':0,'is_off':1})
        else:
            lunch_mins = int(request.form.get(f'd{i}_lunch', 30) or 0)
            shifts.append({
                'day_index':        i,
                'start_time':       request.form.get(f'd{i}_start','09:00') or None,
                'end_time':         request.form.get(f'd{i}_end','17:00')   or None,
                'lunch_start_time': request.form.get(f'd{i}_lstart')        or None,
                'lunch_minutes':    lunch_mins,
                'is_off':           0,
            })
    return shifts

def _save_shifts(conn, request_id, shifts, proposed_by):
    conn.execute("DELETE FROM shifts WHERE request_id=?", (request_id,))
    for s in shifts:
        conn.execute("""
            INSERT INTO shifts
              (request_id,day_index,start_time,end_time,lunch_start_time,lunch_minutes,is_off,proposed_by)
            VALUES (?,?,?,?,?,?,?,?)""",
            (request_id, s['day_index'], s['start_time'], s['end_time'],
             s['lunch_start_time'], s['lunch_minutes'], s['is_off'], proposed_by))

def _load_shifts(conn, request_id):
    return {s['day_index']: dict(s) for s in conn.execute(
        "SELECT * FROM shifts WHERE request_id=? ORDER BY day_index", (request_id,)
    ).fetchall()}

@app.route('/s/<token>', methods=['GET','POST'])
def employee_form(token):
    conn = db()
    req = conn.execute("""
        SELECT r.*, e.name, e.email FROM schedule_requests r
        JOIN employees e ON e.id=r.employee_id
        WHERE r.employee_token=? AND r.status IN ('pending','changes_requested')
    """, (token,)).fetchone()

    if not req:
        conn.close()
        return render_template('error.html',
            msg="This link has already been used or has expired. "
                "Contact your manager if you need to make changes.")

    existing = _load_shifts(conn, req['id'])

    if request.method == 'POST':
        shifts = _parse_shifts_from_form()
        _save_shifts(conn, req['id'], shifts, 'employee')
        new_status = 'employee_approved' if req['status']=='changes_requested' else 'submitted'
        conn.execute("UPDATE schedule_requests SET status=?,submitted_at=datetime('now') WHERE id=?",
                     (new_status, req['id']))
        conn.commit()
        cfg = get_config()
        try:
            _notify_manager(dict(req), shifts, cfg)
        except Exception as e:
            app.logger.error(f"Manager notify failed: {e}")
        conn.close()
        return render_template('submitted.html', name=req['name'],
                               is_resubmit=(req['status']=='changes_requested'))

    week_start = date.fromisoformat(req['week_start'])
    day_data   = build_day_data(week_start, existing)
    conn.close()
    return render_template('employee_form.html',
        req=dict(req), week_start=week_start,
        week_end=week_start+timedelta(days=6),
        day_data=day_data, lunch_options=LUNCH_OPTIONS,
        is_changes=(req['status']=='changes_requested'),
        manager_notes=req['manager_notes'])

@app.route('/a/<token>')
def accept_changes(token):
    conn = db()
    req = conn.execute("""
        SELECT r.*, e.name FROM schedule_requests r
        JOIN employees e ON e.id=r.employee_id
        WHERE r.accept_token=? AND r.status='changes_requested'
    """, (token,)).fetchone()
    if not req:
        conn.close()
        return render_template('error.html', msg="This link is no longer valid.")
    conn.execute("UPDATE schedule_requests SET status='employee_approved',submitted_at=datetime('now') WHERE id=?",
                 (req['id'],))
    conn.commit()
    conn.close()
    return render_template('submitted.html', name=req['name'], is_resubmit=True, quick_accept=True)

# ─── Manager routes ───────────────────────────────────────────────────────────

@app.route('/admin')
@mgr_only
def manager_dashboard():
    conn = db()
    weeks = [r['week_start'] for r in conn.execute(
        "SELECT DISTINCT week_start FROM schedule_requests ORDER BY week_start DESC"
    ).fetchall()]
    sel = request.args.get('week', weeks[0] if weeks else None)
    rows = []
    if sel:
        for req in conn.execute("""
            SELECT r.*, e.name, e.email FROM schedule_requests r
            JOIN employees e ON e.id=r.employee_id
            WHERE r.week_start=? ORDER BY e.name
        """, (sel,)).fetchall():
            shifts = _load_shifts(conn, req['id'])
            total  = sum(worked_hours(s['start_time'],s['end_time'],
                                      s['lunch_minutes'],s['is_off'])
                         for s in shifts.values())
            rows.append({'req': dict(req), 'shifts': shifts,
                         'total': round(total,1),
                         'status_meta': STATUS_META.get(req['status'],
                                                        (req['status'],'secondary'))})
    conn.close()
    return render_template('manager_dashboard.html',
        rows=rows, weeks=weeks, sel=sel,
        week_start=date.fromisoformat(sel) if sel else None,
        days_short=DAYS_SHORT,
        worked_hours=worked_hours)

@app.route('/admin/review/<token>', methods=['GET','POST'])
@mgr_only
def manager_review(token):
    conn = db()
    req = conn.execute("""
        SELECT r.*, e.name, e.email FROM schedule_requests r
        JOIN employees e ON e.id=r.employee_id
        WHERE r.manager_token=?
    """, (token,)).fetchone()
    if not req:
        conn.close()
        abort(404)

    existing = _load_shifts(conn, req['id'])

    if request.method == 'POST':
        action        = request.form.get('action')
        manager_notes = request.form.get('manager_notes','').strip()
        shifts        = _parse_shifts_from_form()
        _save_shifts(conn, req['id'], shifts,
                     'employee' if action=='approve' else 'manager')
        cfg = get_config()

        if action == 'approve':
            conn.execute(
                "UPDATE schedule_requests SET status='approved',approved_at=datetime('now'),manager_notes=? WHERE id=?",
                (manager_notes, req['id']))
            conn.commit()
            try:
                _send_approval(dict(req), shifts, cfg)
            except Exception as e:
                app.logger.error(f"Approval email: {e}")
        else:
            conn.execute(
                "UPDATE schedule_requests SET status='changes_requested',manager_notes=? WHERE id=?",
                (manager_notes, req['id']))
            conn.commit()
            try:
                _send_changes(dict(req), shifts, manager_notes, cfg)
            except Exception as e:
                app.logger.error(f"Changes email: {e}")

        conn.close()
        return redirect(url_for('manager_dashboard', week=req['week_start']))

    week_start = date.fromisoformat(req['week_start'])
    day_data   = build_day_data(week_start, existing)
    total      = sum(worked_hours(d['start_time'],d['end_time'],
                                  d['lunch_minutes'],d['is_off'])
                     for d in day_data if not d['is_off'])
    conn.close()
    return render_template('manager_review.html',
        req=dict(req), week_start=week_start,
        week_end=week_start+timedelta(days=6),
        day_data=day_data, lunch_options=LUNCH_OPTIONS,
        total_hours=round(total,1),
        status_meta=STATUS_META.get(req['status'],(req['status'],'secondary')))

@app.route('/admin/export/<week_start>')
@mgr_only
def export_excel(week_start):
    conn = db()
    reqs = conn.execute("""
        SELECT r.*, e.name FROM schedule_requests r
        JOIN employees e ON e.id=r.employee_id
        WHERE r.week_start=? ORDER BY e.name
    """, (week_start,)).fetchall()

    data = []
    for req in reqs:
        shifts = _load_shifts(conn, req['id'])
        data.append({'name':req['name'],'status':req['status'],'shifts':shifts})
    conn.close()

    wb  = _excel_export(week_start, data)
    buf = BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name=f'schedule_{week_start}.xlsx')

# ─── Excel builder ────────────────────────────────────────────────────────────

def _xfill(h): return PatternFill("solid", fgColor=h)
def _xfont(bold=False, color="000000", size=10):
    return Font(bold=bold, color=color, size=size, name='Calibri')
def _xalign(h='center', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def _xthin():
    s=Side(style='thin',color='BFBFBF')
    return Border(top=s,bottom=s,left=s,right=s)
def _xmed():
    s=Side(style='medium',color='1F4E79')
    return Border(top=s,bottom=s,left=s,right=s)

def _excel_export(week_start_str, employee_data):
    wb  = Workbook()
    wk  = date.fromisoformat(week_start_str)

    # ── Tab 1: Schedule Summary ───────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Schedule Summary"

    ws1.column_dimensions['A'].width = 22
    for c in range(2, 11):
        ws1.column_dimensions[get_column_letter(c)].width = 16

    ws1.row_dimensions[1].height = 34
    ws1.merge_cells('A1:J1')
    c=ws1['A1']
    c.value   = f"Employee Schedule  —  Week of {wk.strftime('%B %d, %Y')}"
    c.fill    = _xfill('1F4E79')
    c.font    = _xfont(bold=True, color='FFFFFF', size=14)
    c.alignment = _xalign()

    ws1.row_dimensions[2].height = 36
    hdrs = ['Employee'] + [
        f"{DAYS_SHORT[i]}\n{(wk+timedelta(days=i)).strftime('%b %d')}"
        for i in range(7)
    ] + ['Status','Total Hrs']
    for ci, h in enumerate(hdrs, 1):
        c = ws1.cell(row=2, column=ci, value=h)
        c.fill      = _xfill('D6E4F0')
        c.font      = _xfont(bold=True, color='1F4E79', size=10)
        c.alignment = _xalign(wrap=True)
        c.border    = _xthin()

    STATUS_COLORS = {'approved':'C6EFCE','employee_approved':'DAEEF3',
                     'submitted':'FFEB9C','changes_requested':'FCE4D6','pending':'F2F2F2'}

    for ri, emp in enumerate(employee_data, 3):
        ws1.row_dimensions[ri].height = 44
        bg = 'FAFAFA' if ri % 2 == 1 else 'FFFFFF'

        c = ws1.cell(row=ri, column=1, value=emp['name'])
        c.font      = _xfont(bold=True, size=11)
        c.alignment = _xalign(h='left')
        c.border    = _xthin()
        c.fill      = _xfill(bg)

        total = 0.0
        for di in range(7):
            s = emp['shifts'].get(di)
            if s:
                hrs = worked_hours(s['start_time'],s['end_time'],s['lunch_minutes'],s['is_off'])
                total += hrs
                if s['is_off']:
                    txt = 'Day Off'; cbg='F2F2F2'
                else:
                    st=s['start_time'] or ''; et=s['end_time'] or ''
                    lm=s['lunch_minutes']
                    ls=s.get('lunch_start_time') or ''
                    lunch_str=(f"{ls}" if ls else "") + (f"  {lm}m lunch" if lm else "  no lunch")
                    txt=f"{st}–{et}\n{lunch_str.strip()}\n{hrs:.1f} hrs"
                    cbg='E2EFDA' if di<5 else 'EDE7F6'
            else:
                txt='—'; cbg='F2F2F2'
            c = ws1.cell(row=ri, column=di+2, value=txt)
            c.fill      = _xfill(cbg)
            c.font      = _xfont(size=9)
            c.alignment = _xalign(wrap=True)
            c.border    = _xthin()

        c = ws1.cell(row=ri, column=9,
                     value=STATUS_META.get(emp['status'],('?',''))[0])
        c.fill      = _xfill(STATUS_COLORS.get(emp['status'],'F2F2F2'))
        c.font      = _xfont(size=9)
        c.alignment = _xalign()
        c.border    = _xthin()

        c = ws1.cell(row=ri, column=10, value=round(total,1))
        c.fill      = _xfill('DAEEF3')
        c.font      = _xfont(bold=True, size=11)
        c.alignment = _xalign()
        c.border    = _xthin()

    ws1.freeze_panes = 'B3'
    ws1.page_setup.orientation='landscape'
    ws1.page_setup.fitToPage=True
    ws1.page_setup.fitToWidth=1

    # ── Tab 2: Visual Timeline (one row per employee per day) ────────
    # Layout: Col A = employee name (merged across 5 day rows),
    #         Col B = day name,  Cols C+ = 30-min time slots 8am–9pm.
    # Scroll vertically to compare employees; left-right shows time of day.
    ws2 = wb.create_sheet("Visual Timeline")

    N_SLOTS  = len(TIMELINE_SLOTS)   # 26 slots (8:00–20:30)
    SLOT_COL = 3                     # time slots start at column C

    ws2.column_dimensions['A'].width = 20   # employee name
    ws2.column_dimensions['B'].width = 12   # day label
    for ci in range(SLOT_COL, SLOT_COL + N_SLOTS):
        ws2.column_dimensions[get_column_letter(ci)].width = 2.6

    # Row 1: title
    last_col_ltr = get_column_letter(SLOT_COL + N_SLOTS - 1)
    ws2.merge_cells(f'A1:{last_col_ltr}1')
    ws2.row_dimensions[1].height = 26
    c = ws2['A1']
    c.value     = (f"Staff Coverage Timeline  —  Week of {wk.strftime('%B %d, %Y')}"
                   "  (Mon–Fri  |  8 AM – 9 PM  |  30-min slots)")
    c.fill      = _xfill('1F4E79')
    c.font      = _xfont(bold=True, color='FFFFFF', size=12)
    c.alignment = _xalign()

    # Row 2: column headers + time slot labels
    ws2.row_dimensions[2].height = 24
    for col, txt in [(1, 'Employee'), (2, 'Day')]:
        c = ws2.cell(row=2, column=col, value=txt)
        c.fill      = _xfill('D6E4F0')
        c.font      = _xfont(bold=True, color='1F4E79', size=10)
        c.alignment = _xalign()
        c.border    = _xthin()

    for si, slot in enumerate(TIMELINE_SLOTS):
        h, m = slot // 60, slot % 60
        label = f"{h}" if m == 0 else f"{m}"   # show hour on-hour, :30 as "30"
        c = ws2.cell(row=2, column=SLOT_COL + si, value=label)
        c.fill      = _xfill('EBF3FB')
        c.font      = Font(size=7, color='1F4E79', name='Calibri', bold=(m == 0))
        c.alignment = Alignment(horizontal='center', vertical='bottom',
                                text_rotation=90)
        c.border    = Border(bottom=Side(style='thin', color='BFBFBF'),
                             left=Side(style='hair', color='DDDDDD'))

    # Data rows — 5 rows per employee (Mon–Fri), then a thin separator row
    cur = 3
    for ei, emp in enumerate(employee_data):
        emp_row_start = cur

        for di in range(5):   # Mon=0 … Fri=4
            ws2.row_dimensions[cur].height = 16
            day_date  = wk + timedelta(days=di)
            day_label = f"{WEEKDAYS_SHORT[di]}  {day_date.strftime('%m/%d')}"
            row_bg    = 'F2F6FA' if ei % 2 == 0 else 'FAFAFA'
            is_last_day = (di == 4)

            c = ws2.cell(row=cur, column=2, value=day_label)
            c.font      = _xfont(size=9, bold=(di == 0))
            c.alignment = _xalign(h='left')
            c.fill      = _xfill(row_bg)
            c.border    = Border(
                bottom=Side(style='medium' if is_last_day else 'hair',
                            color='BFBFBF' if is_last_day else 'E0E0E0'),
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'))

            # Compute this employee's shift for this day
            s = emp['shifts'].get(di)
            if s and not s['is_off']:
                s_min  = to_min(s['start_time'])
                e_min  = to_min(s['end_time'])
                lm     = s.get('lunch_minutes', 0) or 0
                ls_min = lunch_start_minutes(s['start_time'], s['end_time'],
                                             s.get('lunch_start_time'), lm)
                le_min = (ls_min + lm) if ls_min is not None else None
            else:
                s_min = e_min = ls_min = le_min = None

            for si, slot in enumerate(TIMELINE_SLOTS):
                col   = SLOT_COL + si
                color = slot_fill(slot, s_min, e_min, ls_min, le_min)
                cell  = ws2.cell(row=cur, column=col)
                cell.fill   = _xfill(color if color else row_bg)
                cell.border = Border(
                    bottom=Side(style='medium' if is_last_day else 'hair',
                                color='BFBFBF' if is_last_day else 'E0E0E0'),
                    left=Side(style='hair', color='D8D8D8'))
            cur += 1

        # Merge employee name vertically across 5 day rows
        ws2.merge_cells(f'A{emp_row_start}:A{cur - 1}')
        c = ws2.cell(row=emp_row_start, column=1, value=emp['name'])
        c.font      = _xfont(bold=True, size=10)
        c.alignment = Alignment(horizontal='left', vertical='center',
                                wrap_text=False, indent=1)
        c.fill      = _xfill('E8F0F8' if ei % 2 == 0 else 'EBF3FB')
        c.border    = Border(
            top=Side(style='medium', color='1F4E79'),
            bottom=Side(style='medium', color='1F4E79'),
            left=Side(style='medium', color='1F4E79'),
            right=Side(style='thin', color='BFBFBF'))

        # Thin separator row between employees
        ws2.row_dimensions[cur].height = 4
        cur += 1

    # Legend
    leg = cur + 1
    ws2.row_dimensions[leg].height = 16
    ws2.cell(row=leg, column=1, value="Legend:").font = _xfont(bold=True, size=9)
    for i, (color, lbl) in enumerate(
            [('4472C4','On Shift'),('FFEB9C','Lunch Break'),('F2F6FA','Off / Not Working')]):
        c = ws2.cell(row=leg, column=3 + i*2); c.fill=_xfill(color); c.border=_xthin()
        c = ws2.cell(row=leg, column=4 + i*2, value=f"  {lbl}")
        c.font=_xfont(size=9); c.alignment=_xalign(h='left')

    ws2.freeze_panes = 'C3'
    ws2.page_setup.orientation = 'landscape'
    ws2.page_setup.fitToPage  = True
    ws2.page_setup.fitToWidth = 1

    return wb

# ─── Email ────────────────────────────────────────────────────────────────────

def _base_url():
    cfg = get_config()
    if 'base_url' in cfg:
        return cfg['base_url'].rstrip('/')
    try:
        ip = socket.gethostbyname(socket.gethostname())
    except Exception:
        ip = '127.0.0.1'
    return f"http://{ip}:5000"

def _smtp_send(to_email, to_name, subject, html, cfg):
    msg = MIMEMultipart('alternative')
    msg['Subject'] = subject
    msg['From']    = f"{cfg['smtp']['from_name']} <{cfg['smtp']['from_email']}>"
    msg['To']      = f"{to_name} <{to_email}>"
    msg.attach(MIMEText(html, 'html', 'utf-8'))
    with smtplib.SMTP(cfg['smtp']['host'], cfg['smtp']['port']) as s:
        s.ehlo(); s.starttls()
        s.login(cfg['smtp']['username'], cfg['smtp']['password'])
        s.sendmail(cfg['smtp']['from_email'], to_email, msg.as_string())

def _wrap(subtitle, body):
    return f"""<!DOCTYPE html><html lang="en">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
<body style="margin:0;padding:0;background:#f0f4f8;font-family:Calibri,Segoe UI,Arial,sans-serif;">
<table width="100%" cellpadding="0" cellspacing="0" bgcolor="#f0f4f8" style="padding:28px 0;">
<tr><td align="center">
<table width="600" cellpadding="0" cellspacing="0"
       style="background:#fff;border-radius:8px;overflow:hidden;box-shadow:0 2px 12px rgba(0,0,0,.09);">
  <tr><td bgcolor="#1F4E79" style="padding:26px 36px;">
    <p style="margin:0;color:#fff;font-size:19px;font-weight:bold;">Weekly Shift Schedule</p>
    <p style="margin:5px 0 0;color:#9DC3E6;font-size:13px;">{subtitle}</p>
  </td></tr>
  <tr><td style="padding:30px 36px;">{body}</td></tr>
  <tr><td bgcolor="#f7f9fb" style="padding:14px 36px;border-top:1px solid #e8ecf0;text-align:center;">
    <p style="margin:0;color:#bbb;font-size:11px;">Sent by your shift scheduler &bull; Do not reply to this email</p>
  </td></tr>
</table>
</td></tr></table>
</body></html>"""

def _shift_table(shifts):
    LUNCH_LABELS = {v:k for k,v in dict(LUNCH_OPTIONS).items()}
    rows = ''
    total = 0.0
    for i, s in enumerate(shifts):
        if s['is_off']:
            detail='Day Off'; info='—'; hrs=0.0; bg='#F2F2F2'
        else:
            hrs = worked_hours(s['start_time'],s['end_time'],s['lunch_minutes'],s['is_off'])
            total += hrs
            st = s.get('start_time') or ''; et = s.get('end_time') or ''
            lm = s.get('lunch_minutes',0) or 0
            ls = s.get('lunch_start_time') or ''
            lunch_str = (f"{ls} " if ls else '') + (f"{lm} min lunch" if lm else 'no lunch break')
            detail = f"{st} – {et}"; info = f"{lunch_str} &bull; {hrs:.1f} hrs"
            bg = '#E2EFDA' if i < 5 else '#EDE7F6'
        rows += (f"<tr>"
                 f"<td width='110' style='padding:8px 12px;font-weight:bold;color:#1F4E79;font-size:13px;"
                 f"border-bottom:1px solid #eee;'>{DAYS[i]}</td>"
                 f"<td style='padding:8px 12px;background:{bg};font-size:13px;"
                 f"border-bottom:1px solid #eee;'>{detail}</td>"
                 f"<td style='padding:8px 12px;color:#666;font-size:12px;"
                 f"border-bottom:1px solid #eee;'>{info}</td></tr>")
    rows += (f"<tr><td colspan='2' style='padding:9px 12px;font-weight:bold;font-size:13px;"
             f"color:#1F4E79;background:#DAEEF3;'>Total Hours</td>"
             f"<td style='padding:9px 12px;font-weight:bold;font-size:14px;"
             f"color:#1F4E79;background:#DAEEF3;'>{total:.1f} hrs</td></tr>")
    return (f"<table width='100%' cellpadding='0' cellspacing='0'"
            f" style='border-collapse:collapse;margin:16px 0;'>"
            f"<tr bgcolor='#1F4E79'>"
            f"<th style='padding:9px 12px;text-align:left;color:#fff;font-size:12px;'>Day</th>"
            f"<th style='padding:9px 12px;text-align:left;color:#fff;font-size:12px;'>Shift</th>"
            f"<th style='padding:9px 12px;text-align:left;color:#fff;font-size:12px;'>Details</th>"
            f"</tr>{rows}</table>")

def _btn(url, label, bg='#1F4E79'):
    return (f"<table cellpadding='0' cellspacing='0' style='margin:20px 0 0;'><tr>"
            f"<td bgcolor='{bg}' style='border-radius:6px;'>"
            f"<a href='{url}' style='display:inline-block;padding:13px 28px;"
            f"color:#fff;font-size:15px;font-weight:bold;text-decoration:none;'>{label} &rarr;</a>"
            f"</td></tr></table>")

def _send_invite(emp, week_start, token, cfg):
    base = _base_url()
    we   = week_start + timedelta(days=6)
    body = (f"<p style='margin:0 0 14px;font-size:16px;color:#333;'>Hi <strong>{emp['name']}</strong>,</p>"
            f"<p style='margin:0 0 18px;font-size:15px;line-height:1.6;color:#555;'>"
            f"Please submit your preferred shift schedule for the week of "
            f"<strong>{week_start.strftime('%B %d')} – {we.strftime('%B %d, %Y')}</strong>.</p>"
            f"<p style='margin:0 0 22px;font-size:14px;line-height:1.6;color:#777;'>"
            f"Set your start time, finish time, and lunch break for each day. "
            f"Mark any days you're unavailable as Day Off.</p>"
            + _btn(f"{base}/s/{token}", "Submit My Schedule") +
            f"<p style='margin:22px 0 0;font-size:12px;color:#bbb;'>"
            f"Please submit by end of day Thursday.</p>")
    _smtp_send(emp['email'], emp['name'],
        f"Submit Your Schedule — Week of {week_start.strftime('%b %d, %Y')}",
        _wrap(f"Week of {week_start.strftime('%B %d, %Y')}", body), cfg)

def _notify_manager(req, shifts, cfg):
    base = _base_url()
    ws   = date.fromisoformat(req['week_start'])
    body = (f"<p style='margin:0 0 14px;font-size:16px;color:#333;'>"
            f"<strong>{req['name']}</strong> submitted their schedule for "
            f"<strong>{ws.strftime('%B %d, %Y')}</strong>.</p>"
            + _shift_table(shifts)
            + _btn(f"{base}/admin/review/{req['manager_token']}", "Review &amp; Approve"))
    _smtp_send(cfg['manager_email'], cfg.get('manager_name','Manager'),
        f"Schedule Submitted: {req['name']} — {ws.strftime('%b %d')}",
        _wrap(f"New submission from {req['name']}", body), cfg)

def _send_changes(req, shifts, notes, cfg):
    base = _base_url()
    ws   = date.fromisoformat(req['week_start'])
    note_html = ''
    if notes:
        note_html = (f"<table width='100%' cellpadding='0' cellspacing='0' style='margin:14px 0;'><tr>"
                     f"<td style='border-left:4px solid #FFC107;background:#FFF9E6;"
                     f"padding:13px 16px;border-radius:0 6px 6px 0;'>"
                     f"<p style='margin:0;font-size:14px;color:#555;'>"
                     f"<strong>Manager's note:</strong> {notes}</p></td></tr></table>")
    body = (f"<p style='margin:0 0 14px;font-size:16px;color:#333;'>Hi <strong>{req['name']}</strong>,</p>"
            f"<p style='margin:0 0 14px;font-size:15px;line-height:1.6;color:#555;'>"
            f"Your manager reviewed your schedule for "
            f"<strong>{ws.strftime('%B %d, %Y')}</strong> and suggested some changes.</p>"
            + note_html + _shift_table(shifts) +
            f"<p style='margin:16px 0 8px;font-size:14px;color:#555;'>"
            f"You can accept these changes with one click, or open the form to modify further.</p>"
            f"<table cellpadding='0' cellspacing='0'><tr>"
            f"<td bgcolor='#2E7D32' style='border-radius:6px;'>"
            f"<a href='{base}/a/{req['accept_token']}' style='display:inline-block;padding:12px 22px;"
            f"color:#fff;font-size:14px;font-weight:bold;text-decoration:none;'>&#10003; Accept Changes</a></td>"
            f"<td width='10'></td>"
            f"<td bgcolor='#455A64' style='border-radius:6px;'>"
            f"<a href='{base}/s/{req['employee_token']}' style='display:inline-block;padding:12px 22px;"
            f"color:#fff;font-size:14px;font-weight:bold;text-decoration:none;'>&#9998; Review &amp; Modify</a>"
            f"</td></tr></table>")
    _smtp_send(req['email'], req['name'],
        f"Schedule Changes — Please Review — {ws.strftime('%b %d')}",
        _wrap(f"Changes proposed for {ws.strftime('%B %d, %Y')}", body), cfg)

def _send_approval(req, shifts, cfg):
    ws = date.fromisoformat(req['week_start'])
    body = (f"<p style='margin:0 0 14px;font-size:16px;color:#333;'>Hi <strong>{req['name']}</strong>,</p>"
            f"<p style='margin:0 0 14px;font-size:15px;line-height:1.6;color:#555;'>"
            f"Your schedule for <strong>{ws.strftime('%B %d, %Y')}</strong> has been "
            f"<strong style='color:#2E7D32;'>confirmed</strong>. &#10003;</p>"
            + _shift_table(shifts) +
            f"<p style='margin:16px 0 0;font-size:12px;color:#bbb;'>"
            f"Please save this email for your records.</p>")
    _smtp_send(req['email'], req['name'],
        f"Schedule Confirmed — Week of {ws.strftime('%b %d, %Y')}",
        _wrap(f"Confirmed: {ws.strftime('%B %d, %Y')}", body), cfg)

# ─── CLI ─────────────────────────────────────────────────────────────────────

def cli_send_invites(week_str):
    try:
        week_start = date.fromisoformat(week_str)
    except ValueError:
        sys.exit(f"Bad date '{week_str}' — use YYYY-MM-DD.")
    init_db()
    cfg  = get_config()
    conn = db()
    for emp in cfg['employees']:
        conn.execute("INSERT OR IGNORE INTO employees (name,email) VALUES (?,?)",
                     (emp['name'], emp['email']))
        row = conn.execute("SELECT * FROM employees WHERE email=?", (emp['email'],)).fetchone()
        if conn.execute("SELECT id FROM schedule_requests WHERE employee_id=? AND week_start=?",
                        (row['id'], week_str)).fetchone():
            print(f"  skip  {emp['name']} (already sent)")
            continue
        et = secrets.token_urlsafe(32)
        at = secrets.token_urlsafe(32)
        mt = secrets.token_urlsafe(32)
        conn.execute("""INSERT INTO schedule_requests
                        (employee_id,week_start,employee_token,accept_token,manager_token)
                        VALUES (?,?,?,?,?)""",
                     (row['id'], week_str, et, at, mt))
        conn.commit()
        try:
            _send_invite(dict(row) | {'email': emp['email']}, week_start, et, cfg)
            print(f"  sent  {emp['name']} <{emp['email']}>")
        except Exception as e:
            print(f"  ERROR {emp['name']}: {e}")
    conn.close()
    print("\nDone. Open http://localhost:5000/admin to manage schedules.")

if __name__ == '__main__':
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest='cmd')

    sp = sub.add_parser('serve')
    sp.add_argument('--host', default='0.0.0.0')
    sp.add_argument('--port', type=int, default=5000)
    sp.add_argument('--debug', action='store_true')

    ip = sub.add_parser('send-invites')
    ip.add_argument('--week', required=True, help='YYYY-MM-DD (Monday)')

    args = p.parse_args()
    if args.cmd == 'serve':
        init_db()
        print(f"\n  Manager dashboard → http://localhost:{args.port}/admin")
        print("  Press Ctrl+C to stop\n")
        app.run(host=args.host, port=args.port, debug=args.debug)
    elif args.cmd == 'send-invites':
        cli_send_invites(args.week)
    else:
        p.print_help()
